"""
Lógica compartilhada para os relatórios do módulo Controladoria Financeira.

Estratégia de extração:
  - Carga inicial  : keyset pagination via R_E_C_N_O_ (batches de BATCH_SIZE)
  - Sync incremental: date-window (max_data_local - LOOKBACK_DIAS) para capturar
                      registros novos E atualizações (ex: baixas de títulos SE1/SE2)
"""
import csv
import io
import os
from services.database import conectar_financeiro, agora
from services.protheus_readonly import executar_select

BATCH_SIZE    = 5000
DATA_INICIO   = '20240101'

# Janela de resync para capturar atualizações (baixas, alterações de saldo, etc.).
# Cada módulo passa seu próprio lookback_dias para sincronizar(). Este valor global
# é o fallback quando nenhum é fornecido (mantém compatibilidade).
def _lookback_dias(env_especifico=None, default=30):
    """Lê env_especifico, com fallback para SYNC_LOOKBACK_DAYS e depois para default."""
    for chave in filter(None, [env_especifico, 'SYNC_LOOKBACK_DAYS']):
        val = os.environ.get(chave)
        if val:
            try:
                return max(1, int(val))
            except (TypeError, ValueError):
                pass
    return default

LOOKBACK_DIAS = _lookback_dias()


def _s(v):
    """Converte valor para string limpa (sem espaços)."""
    return str(v).strip() if v is not None else ''


def _f(v):
    """Converte valor para float seguro."""
    try:
        return float(v) if v is not None else 0.0
    except (TypeError, ValueError):
        return 0.0


# ─── Cursor keyset ────────────────────────────────────────────────────────────

def obter_cursor(tabela):
    conn = conectar_financeiro()
    try:
        row = conn.execute(
            'SELECT ultimo_recno FROM sync_cursor WHERE tabela = ?', (tabela,)
        ).fetchone()
        return row['ultimo_recno'] if row else 0
    finally:
        conn.close()


def _salvar_cursor(conn, tabela, recno):
    conn.execute(
        '''INSERT INTO sync_cursor (tabela, ultimo_recno, atualizado_em)
           VALUES (?, ?, ?)
           ON CONFLICT(tabela) DO UPDATE SET
               ultimo_recno  = excluded.ultimo_recno,
               atualizado_em = excluded.atualizado_em''',
        (tabela, recno, agora())
    )


# ─── Sync log ─────────────────────────────────────────────────────────────────

def registrar_sync(sync_log_tabela, novos, status, erro_resumo=None):
    conn = conectar_financeiro()
    try:
        conn.execute(
            f'INSERT INTO {sync_log_tabela} (executado_em, registros_novos, status, erro_resumo) '
            f'VALUES (?, ?, ?, ?)',
            (agora(), novos, status, erro_resumo)
        )
        conn.commit()
    finally:
        conn.close()


# ─── Extração em páginas (carga inicial) ──────────────────────────────────────

def limpar_para_refresh(sqlite_tabela: str, nome_cursor: str) -> None:
    """Limpa a tabela de dados e o cursor de paginação para forçar carga completa."""
    conn = conectar_financeiro()
    try:
        conn.execute(f'DELETE FROM {sqlite_tabela}')
        conn.execute('DELETE FROM sync_cursor WHERE tabela = ?', (nome_cursor,))
        conn.commit()
    finally:
        conn.close()


def carga_inicial(nome_cursor, query_paginada, sqlite_tabela, sync_log_tabela, fn_upsert):
    """Executa carga inicial via keyset se a tabela estiver vazia."""
    conn = conectar_financeiro()
    try:
        total_local = conn.execute(f'SELECT COUNT(*) FROM {sqlite_tabela}').fetchone()[0]
    finally:
        conn.close()

    if total_local > 0:
        print(f'[FINANCEIRO] {sqlite_tabela}: dados já existem ({total_local} registros). Pulando carga inicial.')
        return 0

    cursor  = 0
    total   = 0
    try:
        while True:
            linhas = executar_select(query_paginada, (cursor,))
            if not linhas:
                break

            max_recno = max(int(r[0]) for r in linhas)
            conn = conectar_financeiro()
            try:
                fn_upsert(conn, linhas)
                _salvar_cursor(conn, nome_cursor, max_recno)
                conn.commit()
            finally:
                conn.close()

            total  += len(linhas)
            cursor  = max_recno
            print(f'[FINANCEIRO] {sqlite_tabela}: carga inicial — {total} registros inseridos...')

            if len(linhas) < BATCH_SIZE:
                break

        registrar_sync(sync_log_tabela, total, 'sucesso' if total > 0 else 'sem_novos')
        # ANALYZE recomendado pelo SQLite após carga grande: atualiza
        # estatísticas usadas pelo query planner. Roda em segundos para
        # tabelas até ~1M linhas.
        if total > 0:
            conn = conectar_financeiro()
            try:
                conn.execute(f'ANALYZE {sqlite_tabela}')
                conn.commit()
            finally:
                conn.close()
        print(f'[FINANCEIRO] {sqlite_tabela}: carga inicial concluída — {total} registros.')
        return total

    except Exception as exc:
        registrar_sync(sync_log_tabela, total, 'erro', str(exc)[:300])
        raise


# ─── Helpers de UPSERT (centralizado) ─────────────────────────────────────────

def construir_upsert_sql(tabela, colunas_header):
    """Monta um UPSERT que ATUALIZA TODAS as colunas no conflito por recno.

    Antes desta padronização, alguns relatórios usavam `DO NOTHING` (updates
    do Protheus eram silenciosamente descartados) e outros usavam `DO UPDATE`
    apenas em algumas colunas (mudanças nas demais ficavam stale).

    Agora todos os módulos financeiros usam este helper, garantindo que
    qualquer alteração em qualquer coluna (saldo, baixa, juros, descrição,
    tudo) é refletida no SQLite local.
    """
    cols = ', '.join(colunas_header)
    placeholders = ', '.join(['?'] * (len(colunas_header) + 1))
    set_clause = ',\n            '.join(
        [f'{c} = excluded.{c}' for c in colunas_header]
    )
    # WHERE evita UPDATE quando nenhuma coluna mudou — preserva total_changes
    # honesto e poupa I/O em SQLite.
    where_clause = ' OR '.join(
        [f'{c} IS NOT excluded.{c}' for c in colunas_header]
    )
    return f'''
    INSERT INTO {tabela} (recno, {cols})
    VALUES ({placeholders})
    ON CONFLICT(recno) DO UPDATE SET
            {set_clause}
    WHERE   {where_clause}
'''


# ─── Sync incremental (date-window) ──────────────────────────────────────────

def _normalizar_para_diff(v):
    """Normaliza um valor para comparação consistente entre SQLite e Protheus.

    - None -> ''
    - float -> 4 casas decimais (evita 12.34 vs 12.34000001 e 12.0 vs 12)
    - int   -> int (mas se vier como float inteiro, normaliza)
    - str   -> strip()
    """
    if v is None:
        return ''
    if isinstance(v, float):
        if v.is_integer():
            return f'{int(v)}'
        return f'{v:.4f}'
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def _contar_mudancas(conn, sqlite_tabela, linhas):
    """Conta quantas linhas do lote do Protheus realmente diferem do SQLite.

    Premissa: a tupla `linhas[i]` segue o layout `(recno, col1, col2, ...)`
    com exatamente as mesmas colunas (e na mesma ordem) que existem em
    `sqlite_tabela` exceto pela PK auto-increment `id`.

    Retorna (novos, atualizados). "Inalterados" são `len(linhas) - novos - atualizados`.

    Esta contagem reflete mutação real de dados — diferente de "linhas trazidas
    pela query window", que é constante quando o Protheus não muda.
    """
    if not linhas:
        return 0, 0

    recnos = [int(r[0]) for r in linhas]
    placeholder = ','.join(['?'] * len(recnos))
    rows = conn.execute(
        f'SELECT * FROM {sqlite_tabela} WHERE recno IN ({placeholder})',
        recnos
    ).fetchall()

    # rows: cada row é (id, recno, col1, col2, ...). Ignoramos o id (índice 0).
    existentes = {}
    for row in rows:
        # Tupla normalizada de (recno, col1, ...) para comparação
        existentes[row[1]] = tuple(_normalizar_para_diff(v) for v in row[1:])

    novos = 0
    atualizados = 0
    for linha in linhas:
        recno = int(linha[0])
        atual = tuple(_normalizar_para_diff(v) for v in linha)  # (recno, col1, ...)
        if recno not in existentes:
            novos += 1
        elif existentes[recno] != atual:
            atualizados += 1
    return novos, atualizados


def _reconciliar_deletes(conn, sqlite_tabela, campo_data_local, data_corte, recnos_protheus):
    """Remove do SQLite os registros que sumiram do Protheus na janela.

    Premissa: o Protheus, na mesma janela [data_corte, ∞), retornou apenas
    `recnos_protheus`. Qualquer recno local cuja data >= data_corte e que NÃO
    está em `recnos_protheus` foi excluído (D_E_L_E_T_='*'), filtrado por nova
    regra de negócio, ou movido para fora do escopo.

    Retorna a quantidade removida.

    NB: registros fora da janela (com data < data_corte) não são tocados — se
    forem deletados retroativamente no Protheus, ficam stale até alguém aumentar
    o lookback ou rodar uma reconciliação completa.
    """
    if not recnos_protheus:
        # Nenhum recno na janela do Protheus = a janela está vazia OU acabou
        # de ser deletada inteira. Para evitar truncar tudo por engano (ex.:
        # falha transitória de TDS que retorna lista vazia), exigimos pelo
        # menos 1 recno do Protheus para reconciliar. Isso é proteção contra
        # falso-positivo destrutivo.
        return 0

    locais = conn.execute(
        f'SELECT recno FROM {sqlite_tabela} WHERE {campo_data_local} >= ?',
        (data_corte,)
    ).fetchall()
    locais_set = {r['recno'] for r in locais}

    a_remover = locais_set - recnos_protheus
    if not a_remover:
        return 0

    placeholder = ','.join(['?'] * len(a_remover))
    cur = conn.execute(
        f'DELETE FROM {sqlite_tabela} WHERE recno IN ({placeholder})',
        list(a_remover)
    )
    return cur.rowcount or len(a_remover)


def sincronizar(
    nome_cursor, query_paginada, query_sync_window,
    sqlite_tabela, sync_log_tabela, fn_upsert,
    campo_data_local,
    lookback_dias=None,
):
    """
    Incremental: busca registros na janela max_data_local - lookback_dias.
    Captura novos, atualizações (baixas/alterações) E exclusões dentro da janela.
    Se não houver dados locais, faz carga inicial via keyset.

    lookback_dias: dias de janela retroativa. Se None, usa o LOOKBACK_DIAS global
    (controlado por SYNC_LOOKBACK_DAYS). Cada módulo deve passar seu próprio valor
    lido via _lookback_dias() para permitir configuração independente por relatório.

    A métrica reportada em `registrar_sync` é "mutações REAIS" (novos +
    atualizados + removidos), não "linhas processadas". Sem essa distinção,
    o histórico mostraria sempre o mesmo número (igual ao tamanho da janela)
    e induziria o usuário a achar que houve atividade quando não houve.
    """
    conn = conectar_financeiro()
    try:
        row = conn.execute(
            f'SELECT MAX({campo_data_local}) AS max_data FROM {sqlite_tabela}'
        ).fetchone()
        max_data = row['max_data'] if row else None
    finally:
        conn.close()

    if not max_data:
        return carga_inicial(nome_cursor, query_paginada, sqlite_tabela, sync_log_tabela, fn_upsert)

    # Calcula data de corte como string YYYYMMDD (formato Protheus).
    # Cap em hoje: datas de vencimento futuras (dados corrompidos no Protheus, ex: ano 5202)
    # fariam data_corte ir para o futuro, zerando a janela de sync permanentemente.
    from datetime import datetime, timedelta
    hoje = datetime.today()
    try:
        base = datetime.strptime(max_data, '%Y%m%d')
    except ValueError:
        base = hoje
    if base > hoje:
        base = hoje
    efetivo = lookback_dias if lookback_dias is not None else LOOKBACK_DIAS
    data_corte = (base - timedelta(days=efetivo)).strftime('%Y%m%d')

    mutacoes = 0
    novos = atualizados = removidos = 0
    try:
        linhas = executar_select(query_sync_window, (data_corte,))
        recnos_protheus = {int(r[0]) for r in linhas} if linhas else set()

        conn = conectar_financeiro()
        try:
            # 1) Detecta novos + atualizados (sem aplicar ainda)
            if linhas:
                try:
                    novos, atualizados = _contar_mudancas(conn, sqlite_tabela, linhas)
                except Exception as exc_diff:
                    print(f'[FINANCEIRO] {sqlite_tabela}: diff falhou ({exc_diff}); '
                          f'aplicando upsert cego.')
                    novos, atualizados = len(linhas), 0

            # 2) Aplica UPSERT (idempotente; com WHERE no DO UPDATE evita
            #    escritas redundantes mesmo se _contar_mudancas falhar).
            if linhas and (novos + atualizados) > 0:
                fn_upsert(conn, linhas)

            # 3) Reconcilia DELETEs na janela: recnos locais que sumiram do
            #    Protheus dentro do mesmo intervalo de datas são removidos.
            removidos = _reconciliar_deletes(
                conn, sqlite_tabela, campo_data_local, data_corte, recnos_protheus
            )

            # 4) Avança o cursor (mesmo sem mutações, para sincronizar o estado).
            if recnos_protheus:
                _salvar_cursor(conn, nome_cursor, max(recnos_protheus))

            conn.commit()
        finally:
            conn.close()

        mutacoes = novos + atualizados + removidos
        status = 'sucesso' if mutacoes > 0 else 'sem_novos'
        registrar_sync(sync_log_tabela, mutacoes, status)
        return mutacoes

    except Exception as exc:
        registrar_sync(sync_log_tabela, mutacoes, 'erro', str(exc)[:300])
        raise


def sincronizar_cadastro(nome_cursor, query_completa, sqlite_tabela, sync_log_tabela, fn_upsert):
    """Snapshot completo para cadastros sem janela de data (ex.: SA2).

    Busca o universo ativo no Protheus, faz upsert das mudanças e remove
    recnos locais que não voltaram (exclusão lógica no Protheus).
    Lista vazia não apaga o local — protege contra falha transitória.
    """
    try:
        linhas = executar_select(query_completa)
    except Exception as exc:
        registrar_sync(sync_log_tabela, 0, 'erro', str(exc)[:300])
        raise

    recnos_protheus = {int(r[0]) for r in linhas} if linhas else set()
    if not recnos_protheus:
        registrar_sync(sync_log_tabela, 0, 'sem_novos')
        return 0

    conn = conectar_financeiro()
    try:
        novos = atualizados = 0
        try:
            novos, atualizados = _contar_mudancas(conn, sqlite_tabela, linhas)
        except Exception as exc_diff:
            print(f'[FINANCEIRO] {sqlite_tabela}: diff falhou ({exc_diff}); upsert cego.')
            novos, atualizados = len(linhas), 0

        if (novos + atualizados) > 0:
            fn_upsert(conn, linhas)

        locais = {r[0] for r in conn.execute(f'SELECT recno FROM {sqlite_tabela}')}
        a_remover = locais - recnos_protheus
        removidos = 0
        if a_remover:
            placeholder = ','.join(['?'] * len(a_remover))
            cur = conn.execute(
                f'DELETE FROM {sqlite_tabela} WHERE recno IN ({placeholder})',
                list(a_remover),
            )
            removidos = cur.rowcount or len(a_remover)

        _salvar_cursor(conn, nome_cursor, max(recnos_protheus))
        conn.commit()
    finally:
        conn.close()

    mutacoes = novos + atualizados + removidos
    registrar_sync(sync_log_tabela, mutacoes, 'sucesso' if mutacoes > 0 else 'sem_novos')
    return mutacoes


# ─── Info e histórico ─────────────────────────────────────────────────────────

def info_relatorio(sqlite_tabela, sync_log_tabela):
    conn = conectar_financeiro()
    try:
        total = conn.execute(f'SELECT COUNT(*) FROM {sqlite_tabela}').fetchone()[0]
        sync  = conn.execute(
            f'SELECT executado_em, registros_novos, status, erro_resumo '
            f'FROM {sync_log_tabela} ORDER BY id DESC LIMIT 1'
        ).fetchone()
    finally:
        conn.close()

    if not sync:
        return total, None, 'nunca', None
    return total, sync['executado_em'], sync['status'], sync['erro_resumo']


def historico_sync(sync_log_tabela, limit=10, offset=0):
    conn = conectar_financeiro()
    try:
        return conn.execute(
            f'SELECT executado_em, registros_novos, status, erro_resumo '
            f'FROM {sync_log_tabela} '
            f'ORDER BY id DESC LIMIT ? OFFSET ?', (limit, offset)
        ).fetchall()
    finally:
        conn.close()


# ─── Helpers de export ────────────────────────────────────────────────────────

def _construir_query_export(sqlite_tabela, colunas_select, campo_data,
                             data_inicio=None, data_fim=None):
    sql    = f'SELECT {colunas_select} FROM {sqlite_tabela}'
    params = []
    conds  = []
    if data_inicio:
        conds.append(f'{campo_data} >= ?')
        params.append(data_inicio)
    if data_fim:
        conds.append(f'{campo_data} <= ?')
        params.append(data_fim)
    if conds:
        sql += ' WHERE ' + ' AND '.join(conds)
    sql += f' ORDER BY {campo_data} DESC'
    return sql, params


_FETCH_CHUNK = 2000


def gerar_csv(sqlite_tabela, colunas_header, colunas_select, campo_data,
              data_inicio=None, data_fim=None):
    """Gera CSV em modo STREAMING.

    Devolve (iterador_de_strings, callable_que_retorna_total).
    O total é pré-contado via COUNT(*) com os mesmos filtros, garantindo
    auditoria precisa mesmo quando o cliente aborta o download antes do fim.

    Footprint: O(_FETCH_CHUNK) linhas em memória de cada vez, em vez de tudo.
    """
    sql, params = _construir_query_export(
        sqlite_tabela, colunas_select, campo_data, data_inicio, data_fim
    )

    # Pré-conta com os mesmos filtros para auditoria resiliente a abort
    count_sql, count_params = _construir_query_export(
        sqlite_tabela, 'COUNT(*)', campo_data, data_inicio, data_fim
    )
    conn_count = conectar_financeiro()
    try:
        total_pre = conn_count.execute(count_sql, count_params).fetchone()[0]
    finally:
        conn_count.close()

    def gerar():
        conn = conectar_financeiro()
        try:
            buf = io.StringIO()
            writer = csv.writer(buf, delimiter=';')
            writer.writerow(colunas_header)
            yield buf.getvalue()
            buf.seek(0); buf.truncate(0)

            cursor = conn.execute(sql, params)
            while True:
                lote = cursor.fetchmany(_FETCH_CHUNK)
                if not lote:
                    break
                for linha in lote:
                    writer.writerow(
                        [str(v).strip() if v is not None else '' for v in linha]
                    )
                yield buf.getvalue()
                buf.seek(0); buf.truncate(0)
        finally:
            conn.close()

    return gerar(), (lambda: total_pre)


def gerar_excel(sqlite_tabela, colunas_header, colunas_select, campo_data,
                titulo_aba, data_inicio=None, data_fim=None):
    """Gera Excel usando openpyxl em modo write_only (footprint reduzido).

    - write_only: linhas são escritas direto no XML em disco-buffer; nenhum DOM
      mantido em memória além da página atual.
    - Auto-width em dois passes: o primeiro pass lê os dados para calcular
      largura_max; o segundo pass escreve o Excel com as larguras já definidas.
      Isso é necessário porque em write_only o cabeçalho XML (<cols>) é gravado
      na primeira chamada a ws.append() — definir column_dimensions após os
      appends não tem efeito (já foi serializado). O custo extra de dois passes
      em SQLite local é desprezível.
    - Cabeçalho com estilo (negrito, fundo escuro, branco).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.utils import get_column_letter

    sql, params = _construir_query_export(
        sqlite_tabela, colunas_select, campo_data, data_inicio, data_fim
    )

    # ── Passo 1: calcular largura máxima por coluna (sem escrever Excel) ──────
    largura_max = [len(nome) for nome in colunas_header]
    total = 0
    conn = conectar_financeiro()
    try:
        cursor = conn.execute(sql, params)
        while True:
            lote = cursor.fetchmany(_FETCH_CHUNK)
            if not lote:
                break
            for linha in lote:
                for idx, val in enumerate(linha):
                    n = len(str(val).strip() if val is not None else '')
                    if n > largura_max[idx]:
                        largura_max[idx] = n
            total += len(lote)
    finally:
        conn.close()

    # ── Passo 2: montar o Excel com column_dimensions definidos ANTES do
    #    primeiro ws.append() para que sejam incluídos no <cols> do XML ────────
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=titulo_aba[:31])

    for i, n in enumerate(largura_max, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(n + 2, 40)

    header_font  = Font(bold=True, color='FFFFFF')
    header_fill  = PatternFill('solid', fgColor='1A1A1A')
    header_align = Alignment(horizontal='center')

    header_cells = []
    for nome in colunas_header:
        cell = WriteOnlyCell(ws, value=nome)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = header_align
        header_cells.append(cell)
    ws.append(header_cells)

    conn = conectar_financeiro()
    try:
        cursor = conn.execute(sql, params)
        while True:
            lote = cursor.fetchmany(_FETCH_CHUNK)
            if not lote:
                break
            for linha in lote:
                ws.append([str(v).strip() if v is not None else '' for v in linha])
    finally:
        conn.close()

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue(), total
